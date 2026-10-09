"""Bounded table extraction into LX records, without database mutations."""

from __future__ import annotations

import csv
import io
import math
from collections.abc import Iterable, Mapping
from datetime import UTC, date, datetime
from pathlib import Path
from uuid import uuid5
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from xlrd.biffh import XLRDError

from lx_dtypes.models.contracts.dtypes_record_persistence import (
    DtypesRecordClassificationChoicePayload,
    DtypesRecordClassificationGroupPayload,
    DtypesRecordFindingPayload,
    DtypesRecordPersistencePayload,
)
from lx_dtypes.models.contracts.table_import import (
    TableColumn,
    TableExtractionError,
    TableExtractor,
    TableMapping,
    TableRecord,
)
from lx_dtypes.models.interface.KnowledgeBase import KnowledgeBase
from lx_dtypes.models.ledger.p_examination.Pydantic import PExamination
from lx_dtypes.models.meta.ReportMeta import ReportPatientInfo
from lx_dtypes.terminology.terminology_loader import load_module_kb
from lx_dtypes.utils.study_setup_yaml import MAX_SETUP_BYTES, parse_setup_yaml
from pydantic import ValidationError

MAX_TABLE_BYTES = 20 * 1024 * 1024
MAX_TABLE_ROWS = 10_000
MAX_TABLE_COLUMNS = 500
MAX_CELL_CHARS = 100_000


def load_table_mapping(mapping_file: Path) -> TableMapping:
    """The trusted local caller supplies a path; HTTP callers use owned mapping IDs."""
    with mapping_file.open("rb") as source:
        raw = source.read(MAX_SETUP_BYTES + 1)
    return parse_table_mapping(raw.decode("utf-8"))


def parse_table_mapping(raw: str) -> TableMapping:
    return TableMapping.model_validate(parse_setup_yaml(raw))


def spacy_patient_extractor(text: str) -> Mapping[str, object]:
    # PatientDataExtractor retains SensitiveMeta state. A fresh instance is
    # mandatory for each input so an unmatched row cannot inherit a prior name.
    from lx_anonymizer.ner.spacy_extractor import PatientDataExtractor

    return dict(PatientDataExtractor()(text))


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return (
            value.date().isoformat()
            if value.time().isoformat() == "00:00:00"
            else value.isoformat()
        )
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("Non-finite numeric cell")
        if value.is_integer():
            value = int(value)
    if isinstance(value, (str, int, float)):
        result = str(value).strip()
        if len(result) > MAX_CELL_CHARS or "\x00" in result:
            raise ValueError("Cell exceeds limits or contains a null character")
        return result
    raise ValueError("Unsupported spreadsheet cell type")


def _table_rows(content: bytes, suffix: str, spec: TableMapping) -> Iterable[list[str]]:
    try:
        yield from _read_table_rows(content, suffix, spec)
    except (csv.Error, BadZipFile, XLRDError, ParseError, KeyError, OSError) as error:
        raise ValueError("Malformed table document") from error


def _read_table_rows(
    content: bytes, suffix: str, spec: TableMapping
) -> Iterable[list[str]]:
    if len(content) > MAX_TABLE_BYTES:
        raise ValueError("Table exceeds 20 MiB")
    if suffix == ".csv":
        reader = csv.reader(
            io.StringIO(content.decode(spec.source.encoding), newline=""),
            delimiter=spec.source.delimiter,
            strict=True,
        )
        for row in reader:
            yield [_cell(value) for value in row]
    elif suffix == ".xlsx":
        from openpyxl import load_workbook

        with ZipFile(io.BytesIO(content)) as archive:
            if (
                sum(member.file_size for member in archive.infolist())
                > MAX_TABLE_BYTES * 5
            ):
                raise ValueError("Expanded workbook exceeds 100 MiB")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        try:
            if spec.source.sheet is None and len(workbook.sheetnames) != 1:
                raise ValueError("Select a sheet for workbooks with multiple sheets")
            if spec.source.sheet and spec.source.sheet not in workbook.sheetnames:
                raise ValueError("Selected sheet is missing")
            sheet = workbook[spec.source.sheet or workbook.sheetnames[0]]
            if sheet.max_column and sheet.max_column > MAX_TABLE_COLUMNS:
                raise ValueError("Workbook exceeds column limit")
            for cells in sheet.iter_rows():
                if any(cell.data_type in {"f", "e"} for cell in cells):
                    raise ValueError(
                        "Formula and error cells are not accepted; export values"
                    )
                yield [_cell(cell.value) for cell in cells]
        finally:
            workbook.close()
    elif suffix == ".xls":
        import xlrd

        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
        try:
            if spec.source.sheet is None and workbook.nsheets != 1:
                raise ValueError("Select a sheet for workbooks with multiple sheets")
            sheet = (
                workbook.sheet_by_name(spec.source.sheet)
                if spec.source.sheet
                else workbook.sheet_by_index(0)
            )
            if (
                sheet.ncols > MAX_TABLE_COLUMNS
                or sheet.nrows > MAX_TABLE_ROWS + spec.source.header_row
            ):
                raise ValueError("Workbook exceeds table limits")
            for number in range(sheet.nrows):
                row: list[str] = []
                for cell in sheet.row(number):
                    if cell.ctype == xlrd.XL_CELL_ERROR:
                        raise ValueError("Spreadsheet error cells are not accepted")
                    value: object = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        if not isinstance(value, float):
                            raise ValueError("Invalid spreadsheet date cell")
                        value = xlrd.xldate_as_datetime(value, workbook.datemode)
                    row.append(_cell(value))
                yield row
        finally:
            workbook.release_resources()
    else:
        raise ValueError("Supported table formats are .csv, .xlsx and .xls")


def _columns(spec: TableMapping) -> dict[str, TableColumn]:
    columns = {
        "row_id": spec.row_id,
        "first_name": spec.first_name,
        "last_name": spec.last_name,
        "dob": spec.dob,
        "gender": spec.gender,
        "examination_date": spec.examination_date,
        "examination": spec.examination,
    }
    for name, group in spec.finding_groups().items():
        columns[f"findings.{name}"] = group.finding
        columns.update(
            {
                f"findings.{name}.{key}": value
                for key, value in group.classifications.items()
            }
        )
    return columns


def _extract(
    column: TableColumn,
    row: Mapping[str, str],
    *,
    target: str,
    number: int,
    extractors: Mapping[str, TableExtractor],
    cache: dict[tuple[str, str], Mapping[str, object]],
) -> str:
    value = (
        column.value
        if column.value is not None
        else column.separator.join(row[name] for name in column.columns).strip()
    )
    if column.extractor is not None:
        key = (column.extractor, value)
        if key not in cache:
            cache[key] = extractors[column.extractor](value)
        extracted = cache[key].get(column.output or "")
        value = _cell(extracted)
    if not value:
        if column.required:
            raise TableExtractionError(number, target, "required_value_missing")
        return ""
    if column.values:
        if value not in column.values:
            raise TableExtractionError(number, target, "unmapped_value")
        value = column.values[value]
    if column.date_format:
        # Native Excel dates are normalized to ISO; text also accepts the
        # explicitly configured source format.
        try:
            parsed_date = date.fromisoformat(value)
        except ValueError:
            parsed_date = datetime.strptime(value, column.date_format).date()
        value = parsed_date.isoformat()
    return value


def parse_table_records(
    mapping_file: Path,
    source_file: Path,
    *,
    extractors: Mapping[str, TableExtractor] | None = None,
) -> list[TableRecord]:
    """Public file API: YAML + CSV/XLS/XLSX → validated LX records, no writes."""
    with source_file.open("rb") as source:
        content = source.read(MAX_TABLE_BYTES + 1)
    return extract_table_records(
        load_table_mapping(mapping_file),
        content,
        source_file.suffix.lower(),
        extractors=extractors,
    )


def extract_table_records(
    spec: TableMapping,
    content: bytes,
    suffix: str,
    *,
    extractors: Mapping[str, TableExtractor] | None = None,
) -> list[TableRecord]:
    registered: dict[str, TableExtractor] = {"spacy_patient": spacy_patient_extractor}
    if extractors is not None:
        registered.update(extractors)
    columns = _columns(spec)
    if any(
        col.extractor is not None and col.extractor not in registered
        for col in columns.values()
    ):
        raise ValueError("Mapping requests an unregistered extractor")
    kb = load_module_kb(
        spec.knowledge_base.knowledge_base_module,
        version=spec.knowledge_base.knowledge_base_version,
    )
    records: list[TableRecord] = []
    seen: set[str] = set()
    header: list[str] | None = None
    for number, values in enumerate(_table_rows(content, suffix, spec), start=1):
        if number < spec.source.header_row:
            continue
        if (
            len(values) > MAX_TABLE_COLUMNS
            or number > MAX_TABLE_ROWS + spec.source.header_row
        ):
            raise TableExtractionError(number, "table", "table_limit_exceeded")
        if header is None:
            header = values
            if (
                not header
                or any(not name for name in header)
                or len(header) != len(set(header))
            ):
                raise TableExtractionError(
                    number, "header", "empty_or_duplicate_header"
                )
            for target, column in columns.items():
                if not set(column.columns).issubset(header):
                    raise TableExtractionError(number, target, "source_column_missing")
            continue
        if not any(values):
            continue
        if len(values) != len(header):
            raise TableExtractionError(number, "table", "row_width_mismatch")
        row = dict(zip(header, values, strict=True))
        extracted: dict[str, str] = {}
        cache: dict[tuple[str, str], Mapping[str, object]] = {}
        for target, column in columns.items():
            try:
                extracted[target] = _extract(
                    column,
                    row,
                    target=target,
                    number=number,
                    extractors=registered,
                    cache=cache,
                )
            except TableExtractionError:
                raise
            except (ValueError, TypeError) as error:
                raise TableExtractionError(
                    number, target, "invalid_extracted_value"
                ) from error
        if extracted["row_id"] in seen:
            raise TableExtractionError(number, "row_id", "duplicate_source_identity")
        seen.add(extracted["row_id"])
        try:
            records.append(_lx_record(spec, extracted, number, kb))
        except (ValueError, KeyError, ValidationError) as error:
            raise TableExtractionError(
                number, "record", "invalid_identity_or_terminology"
            ) from error
    if not records:
        raise ValueError("Table contains no data rows")
    return records


def _lx_record(
    spec: TableMapping, values: Mapping[str, str], number: int, kb: KnowledgeBase
) -> TableRecord:
    source_id = uuid5(spec.mapping_id, values["row_id"])
    examination_date = date.fromisoformat(values["examination_date"])
    timestamp = datetime.combine(examination_date, datetime.min.time(), tzinfo=UTC)
    examination = values["examination"]
    kb.get_examination(examination)
    findings: list[DtypesRecordFindingPayload] = []
    for group_name, group in spec.finding_groups().items():
        finding_name = values[f"findings.{group_name}"]
        if not finding_name:
            if any(
                values[f"findings.{group_name}.{key}"] for key in group.classifications
            ):
                raise ValueError("Classification without finding")
            continue
        finding_id = uuid5(source_id, group_name)
        choices = [
            DtypesRecordClassificationChoicePayload(
                uuid=uuid5(finding_id, classification),
                created_at=timestamp,
                classification=classification,
                classification_choice=values[f"findings.{group_name}.{classification}"],
                patient_finding_classifications=str(finding_id),
            )
            for classification in group.classifications
            if values[f"findings.{group_name}.{classification}"]
        ]
        findings.append(
            DtypesRecordFindingPayload(
                uuid=finding_id,
                created_at=timestamp,
                finding=finding_name,
                patient_examination=examination,
                patient_finding_classifications=[
                    DtypesRecordClassificationGroupPayload(
                        uuid=finding_id,
                        created_at=timestamp,
                        patient_finding=str(finding_id),
                        patient_finding_classification_choices=choices,
                    )
                ],
            )
        )
    record = DtypesRecordPersistencePayload(
        uuid=source_id,
        created_at=timestamp,
        patient=str(source_id),
        date=timestamp,
        examination=examination,
        patient_findings=findings,
        knowledge_base_module=spec.knowledge_base.knowledge_base_module,
        knowledge_base_version=spec.knowledge_base.knowledge_base_version,
    )
    kb.assert_examination_admissibility(
        PExamination.model_validate(record.model_dump(mode="python"))
    )
    return TableRecord(
        row_number=number,
        source_id=source_id,
        examination_date=examination_date,
        patient=ReportPatientInfo(
            first_name=values["first_name"],
            last_name=values["last_name"],
            dob=date.fromisoformat(values["dob"]),
            gender=values["gender"],
        ),
        record=record,
    )
