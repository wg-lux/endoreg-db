from datetime import UTC, datetime
from io import BytesIO
from typing import cast
from unittest.mock import patch

import pytest
from django.core.exceptions import ValidationError
from openpyxl import load_workbook
from rest_framework.test import APIClient
from lx_dtypes.models.contracts.dtypes_record_persistence import (
    DtypesRecordPersistencePayload,
)
from lx_dtypes.models.contracts.examination_revision import (
    CaseExaminationSelection,
    ExaminationRevisionCreate,
    ExaminationRevisionSnapshot,
    ExaminationRevisionUpdate,
    ExaminationEvidenceSelection,
)
from lx_dtypes.models.contracts.json_types import JsonObject
from lx_dtypes.models.ledger.medical.Write import PatientLabValueCreate

from endoreg_db.models import (
    Case,
    Examination,
    Patient,
    PatientExamination,
    LabValue,
    Unit,
    Center,
)
from endoreg_db.models.report.patient_examination_report import PatientExaminationReport
from endoreg_db.models.media.video.video_file import VideoFile
from endoreg_db.models.medical.patient.examination_revision import ExaminationRevision
from endoreg_db.services.cases.examination_revisions import (
    capture_examination_revision,
    delete_examination_revision,
    finalize_examination_revision,
    select_case_examination_revision,
    update_examination_revision,
)
from endoreg_db.services.cases.lifecycle import close_case, reopen_case
from endoreg_db.services.studies.revision_export import (
    case_export_workbook,
    selected_case_export,
)


def examination_case() -> tuple[PatientExamination, Case]:
    patient = Patient.objects.create(
        first_name="Synthetic",
        last_name="Research",
        patient_hash="revision-test",
        is_real_person=False,
    )
    examination = Examination.objects.create(name="revision_exam")
    instance = PatientExamination.objects.create(
        patient=patient, examination=examination, hash="revision-exam"
    )
    instance.dtypes_record = DtypesRecordPersistencePayload(
        patient=str(patient.pk), examination=examination.name
    ).model_dump(mode="json")
    instance.save()
    case = Case.objects.create(
        patient=patient, start_date=datetime(2026, 1, 1, tzinfo=UTC)
    )
    case.patient_examinations.add(instance)
    return instance, case


@pytest.mark.django_db
def test_persisted_draft_crud_selection_snapshot_export_and_follow_up() -> None:
    examination, case = examination_case()
    unit = Unit.objects.create(name="mg", abbreviation="mg")
    LabValue.objects.create(name="example", default_unit=unit)
    first = capture_examination_revision(
        examination,
        ExaminationRevisionCreate(
            patient_examination_id=examination.pk, title="Excel-only draft"
        ),
    )
    alternative = capture_examination_revision(
        examination,
        ExaminationRevisionCreate(
            patient_examination_id=examination.pk, title="Alternative"
        ),
    )
    first.refresh_from_db()
    snapshot = ExaminationRevisionSnapshot.model_validate(first.snapshot)
    snapshot.content.lab_values.append(
        PatientLabValueCreate.model_validate(
            {
                "lab_value": "example",
                "value": 0.0,
                "timestamp": "2026-01-01T10:00:00Z",
                "unit": "mg",
            }
        )
    )
    first = update_examination_revision(
        first,
        ExaminationRevisionUpdate(
            expected_version=1, title="Complete draft", content=snapshot.content
        ),
    )
    with pytest.raises(ValueError, match="changed"):
        update_examination_revision(
            first,
            ExaminationRevisionUpdate(
                expected_version=1, title="Stale", content=snapshot.content
            ),
        )
    with pytest.raises(ValueError, match="Finalize"):
        select_case_examination_revision(
            case, CaseExaminationSelection(revision_id=first.pk, expected_version=2)
        )
    with pytest.raises(ValueError, match="Select one"):
        close_case(instance=case)
    first = finalize_examination_revision(first, 2)
    select_case_examination_revision(
        case, CaseExaminationSelection(revision_id=first.pk, expected_version=2)
    )
    # Current live edits do not rewrite a finalized snapshot or add the unselected alternative.
    examination.dtypes_record["patient_findings"] = []
    examination.save()
    export = selected_case_export(case)
    records = export["examinations"]
    assert isinstance(records, list) and len(records) == 1
    workbook = load_workbook(BytesIO(case_export_workbook(export)))
    assert workbook["Laboratory"].max_row == 2
    assert workbook["Laboratory"].cell(2, 5).value == 0
    assert workbook["Laboratory"].cell(2, 5).data_type == "n"
    assert workbook["Clinical values"].max_row > 5
    assert workbook["Evidence"].max_row > 1
    first.title = "Mutated final"
    with pytest.raises(ValidationError, match="immutable"):
        first.save()
    first.refresh_from_db()
    with pytest.raises(ValueError, match="retained"):
        delete_examination_revision(first, 2)
    delete_examination_revision(alternative, 1)
    assert not ExaminationRevision.objects.filter(pk=alternative.pk).exists()
    closed = close_case(instance=case, end_date=datetime(2026, 1, 2, tzinfo=UTC))
    assert closed.is_closed
    reopened = reopen_case(instance=closed)
    assert reopened.is_active and reopened.end_date is None
    followup = PatientExamination.objects.create(
        patient=examination.patient,
        examination=examination.examination,
        hash="followup",
    )
    reopened.patient_examinations.add(followup)
    with pytest.raises(ValueError, match="Select one"):
        selected_case_export(reopened)


@pytest.mark.django_db
def test_revision_api_crud_and_foreign_owner_rejection(api_client: APIClient) -> None:
    examination, case = examination_case()
    url = "/api/examination-revisions/"
    assert api_client.get(url).status_code == 422
    assert api_client.get(url, {"patient_examination_id": "invalid"}).status_code == 422
    response = api_client.post(
        url, {"patient_examination_id": examination.pk, "title": "Draft"}, format="json"
    )
    assert response.status_code == 201, response.content
    data = response.json()
    record_url = f"{url}{data['id']}/"
    assert api_client.get(record_url).status_code == 200
    content = data["snapshot"]["content"]
    content["examination"]["patient"] = "99999999"
    response = api_client.put(
        record_url,
        {"expected_version": 1, "title": "Foreign", "content": content},
        format="json",
    )
    assert response.status_code == 409, response.content
    assert (
        api_client.post(
            f"{record_url}finalize/", {"expected_version": 1}, format="json"
        ).status_code
        == 200
    )
    response = api_client.post(
        f"/api/cases/{case.case_id}/select-examination/",
        {"revision_id": data["id"], "expected_version": 1},
        format="json",
    )
    assert response.status_code == 200, response.content
    response = api_client.get(f"/api/cases/{case.case_id}/research-export/?output=xlsx")
    assert response.status_code == 200, response.content
    assert load_workbook(BytesIO(response.content))["Examinations"].max_row == 2
    assert (
        api_client.delete(
            record_url, {"expected_version": 1}, format="json"
        ).status_code
        == 409
    )
    response = api_client.post(
        "/api/cases/research-export/",
        {"case_ids": [str(case.case_id)], "output": "json"},
        format="json",
    )
    assert response.status_code == 200, response.content
    assert response.json()["cases"][0]["examinations"][0]["revision_id"] == data["id"]
    assert response["Cache-Control"] == "private, no-store"
    with patch(
        "endoreg_db.services.studies.revision_export.case_export_workbook",
        side_effect=ValueError(
            "Content cannot be represented losslessly in XLSX; use JSON export"
        ),
    ):
        single = api_client.get(
            f"/api/cases/{case.case_id}/research-export/?output=xlsx"
        )
        batch = api_client.post(
            "/api/cases/research-export/",
            {"case_ids": [str(case.case_id)], "output": "xlsx"},
            format="json",
        )
    for failed in (single, batch):
        assert failed.status_code == 409
        assert "use JSON export" in failed.json()["detail"]
    # A corrupt persisted snapshot must never cross the patient boundary on export.
    snapshot = data["snapshot"]
    snapshot["content"]["examination"]["patient"] = "99999999"
    ExaminationRevision.objects.filter(pk=data["id"]).update(snapshot=snapshot)
    for output in ("json", "xlsx"):
        failed = api_client.get(
            f"/api/cases/{case.case_id}/research-export/?output={output}"
        )
        assert failed.status_code == 409
        assert "ownership" in failed.json()["detail"]


@pytest.mark.django_db
@pytest.mark.parametrize(
    "with_report,with_video", [(True, False), (False, True), (True, True)]
)
def test_optional_evidence_is_frozen_and_changed_media_cannot_finalize(
    with_report: bool, with_video: bool
) -> None:
    examination, _ = examination_case()
    report = (
        PatientExaminationReport.objects.create(
            patient_examination=examination,
            template_name="example",
            template_version="1",
            rendered_text="Synthetic report",
            dtypes_record=examination.dtypes_record,
        )
        if with_report
        else None
    )
    video = (
        VideoFile.objects.create(
            examination=examination,
            processed_video_hash="a" * 64,
            center=Center.objects.create(name="revision-center"),
        )
        if with_video
        else None
    )
    draft = capture_examination_revision(
        examination,
        ExaminationRevisionCreate(
            patient_examination_id=examination.pk,
            title="Evidence selection",
            selection=ExaminationEvidenceSelection(
                report_id=report.pk if report else None,
                video_id=video.pk if video else None,
            ),
        ),
    )
    if report:
        report.rendered_text = "Later report edit"
        report.save()
        assert (
            ExaminationRevisionSnapshot.model_validate(
                draft.snapshot
            ).evidence.report_text
            == "Synthetic report"
        )
    # Media byte integrity is covered by the media export tests; verify this service
    # delegates to that gate and independently pins the captured generation.
    with patch(
        "endoreg_db.services.cases.examination_revisions.validate_video_download"
    ) as validate:
        if video:
            VideoFile.objects.filter(pk=video.pk).update(processed_video_hash="b" * 64)
            with pytest.raises(ValueError, match="generation changed"):
                finalize_examination_revision(draft, 1)
            draft.refresh_from_db()
            assert draft.finalized_at is None
            VideoFile.objects.filter(pk=video.pk).update(processed_video_hash="a" * 64)
        final = finalize_examination_revision(draft, 1)
        assert final.finalized_at is not None
        assert validate.call_count == (2 if video else 0)
    copied = capture_examination_revision(
        examination,
        ExaminationRevisionCreate(
            patient_examination_id=examination.pk,
            title="New alternative",
            source_revision_id=final.pk,
        ),
    )
    assert copied.finalized_at is None and copied.snapshot == final.snapshot
    if video:
        VideoFile.objects.filter(pk=video.pk).update(examination=None)
        with pytest.raises(ValueError, match="no longer available"):
            finalize_examination_revision(copied, 1)
    other_exam = PatientExamination.objects.create(
        patient=examination.patient,
        examination=examination.examination,
        hash="other-same-type",
    )
    final.patient_examination = other_exam
    with pytest.raises(ValidationError, match="immutable"):
        final.save()


def _workbook_payload() -> JsonObject:
    snapshot = ExaminationRevisionSnapshot.model_validate(
        {
            "content": {"examination": {"patient": "1", "examination": "=unsafe"}},
            "evidence": {},
        }
    )
    payload: JsonObject = {
        "schema_version": "1.0",
        "generated_at": "2026-01-01T10:00:00Z",
        "case_id": "0b5c337e-3a39-4138-a911-8ba1cb95360a",
        "patient_pseudonym": "pseudo",
        "center_key": None,
        "is_closed": False,
        "examinations": [
            {
                "patient_examination_id": 1,
                "revision_id": 2,
                "revision_version": 1,
                "finalized_at": "2026-01-01T10:00:00Z",
                "snapshot": cast(JsonObject, snapshot.model_dump(mode="json")),
            }
        ],
    }
    return payload


def test_workbook_escapes_formulas_and_preserves_null_and_repeated_values() -> None:
    payload = _workbook_payload()
    records = payload["examinations"]
    assert isinstance(records, list) and isinstance(records[0], dict)
    snapshot = records[0]["snapshot"]
    assert isinstance(snapshot, dict) and isinstance(snapshot["evidence"], dict)
    # Narrative is deliberately excluded from the statistical workbook.
    snapshot["evidence"]["report_text"] = "x" * 32768
    workbook = load_workbook(BytesIO(case_export_workbook(payload)))
    assert workbook["Examinations"].cell(2, 5).data_type == "s"
    assert workbook["Examinations"].cell(2, 5).value == "=unsafe"
    rows = list(workbook["Clinical values"].values)
    assert any(row[1] == "/lab_values" and row[3] == "empty_list" for row in rows)
    repeated = load_workbook(
        BytesIO(
            case_export_workbook(
                [
                    payload,
                    {**payload, "case_id": "2b5c337e-3a39-4138-a911-8ba1cb95360a"},
                ]
            )
        )
    )
    assert repeated["Cases"].max_row == 3
    assert repeated["Case examinations"].max_row == 3
    assert repeated["Examinations"].max_row == 2
    entries = payload["examinations"]
    assert isinstance(entries, list) and isinstance(entries[0], dict)
    entries[0]["revision_id"] = 9007199254740993
    exact = load_workbook(BytesIO(case_export_workbook(payload)))
    assert exact["Examinations"].cell(2, 2).value == "9007199254740993"


def test_workbook_rejects_ambiguous_identity_and_unfinalized_content() -> None:
    from copy import deepcopy

    payload = _workbook_payload()
    with pytest.raises(ValueError, match="Case identities"):
        case_export_workbook([payload, payload])
    conflicting = deepcopy(payload)
    conflicting["case_id"] = "2b5c337e-3a39-4138-a911-8ba1cb95360a"
    records = conflicting["examinations"]
    assert isinstance(records, list) and isinstance(records[0], dict)
    records[0]["revision_version"] = 2
    with pytest.raises(ValueError, match="Conflicting content"):
        case_export_workbook([payload, conflicting])
    records[0]["finalized_at"] = None
    with pytest.raises(ValueError):
        case_export_workbook(conflicting)


@pytest.mark.parametrize("text", ["x" * 32768, "control\x00character"])
def test_workbook_rejects_content_excel_would_truncate_or_corrupt(text: str) -> None:
    payload = _workbook_payload()
    payload["patient_pseudonym"] = text
    with pytest.raises(ValueError, match="use JSON export"):
        case_export_workbook(payload)
